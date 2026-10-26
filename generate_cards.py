#!/usr/bin/env python3
"""ELECTOR card generation: Excel → PNG/PDF cards + optional sticker sheets."""

from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas as pdf_canvas

ROOT = Path(__file__).resolve().parent
TEMPLATES = ROOT / "ELECTOR" / "cardtemps"
ICONS = ROOT / "ELECTOR" / "iconslabels"
AVATAR_DIR = ROOT / "ELECTOR" / "avatar"
DEFAULT_AVATAR = AVATAR_DIR / "default-img.png"
DEFAULT_XLSX = ROOT / "ELECTOR.xlsx"
STICKER_TEMPLATE = ROOT / "Sticker-temp.png"
OUTPUT_DIR = ROOT / "output"
FONTS_DIR = ROOT / "fonts"
FONT_REGULAR = FONTS_DIR / "Arial.ttf"
FONT_BOLD = FONTS_DIR / "Arial Bold.ttf"

# Poker card 2.5"×3.5" @ 400 DPI → 1000×1400
DPI = 400
CARD_INCH_W, CARD_INCH_H = 2.5, 3.5
CARD_W = int(round(CARD_INCH_W * DPI))
CARD_H = int(round(CARD_INCH_H * DPI))

# Sticker sheet: letter @ 300 DPI, 3×3 of 60×85 mm cards
STICKER_DPI = 300
STICKER_PER_SHEET = 9
STICKER_CARD_W_PX = int(round(60.0 / 25.4 * STICKER_DPI))  # ~709
STICKER_CARD_H_PX = int(round(85.0 / 25.4 * STICKER_DPI))  # ~1004
STICKER_SLOTS = [
    (158, 94), (927, 94), (1695, 94),
    (158, 1159), (927, 1159), (1695, 1159),
    (158, 2224), (927, 2224), (1695, 2224),
]

_DESIGN_W = 1500
SCALE = CARD_W / _DESIGN_W


def sx(v: float) -> int:
    return int(round(v * SCALE))


def pt_px(pt: float) -> int:
    return max(1, int(round(pt * DPI / 72)))


PARTY_TEMPLATE_PREFIX = {
    "DEMOCRAT": "dem",
    "REPUBLICAN": "rep",
    "GREEN": "green",
    "LIBERTARIAN": "lib",
    "INDEPENDENT": "ind",
    "UNIVERSAL": "uni",
}

PARTY_SHORT = {
    "DEMOCRAT": "DEM",
    "REPUBLICAN": "REP",
    "GREEN": "GRN",
    "LIBERTARIAN": "LIB",
    "INDEPENDENT": "IND",
    "UNIVERSAL": "UNI",
}

TYPE_SUFFIX = {"ALLY": "ally", "CANDIDATE": "can", "EFFECT": "eff"}

# Shared footer / party chrome for candidate + effect templates
# Keys match ELECTOR.xlsx column headings where applicable.
_FOOTER = {
    "PARTY": (1216, 1778, 1381, 1945),          # party emblem circle
    "PARTYABB": (1160, 1994, 1440, 2040),
    "FOOTER_CLEAR": (442, 1970, 770, 2060),      # layout-only wipe behind SET/ID
    "SET": (442, 1970, 560, 2018),
    "YEAR": (442, 2022, 760, 2052),              # copyright line uses YEAR
    "ID": (800, 1968, 1120, 2004),
    "CARDNUM": (800, 2000, 1120, 2044),          # print run CARDNUM/SETTOTAL
}
_FOOTER_FONTS = {
    "SET": 5.5,
    "YEAR": 2.3,
    "ID": 3.78,
    "CARDNUM": 4.73,
}

# Layouts for 1500×2100 design canvas; scaled via sx() / pt_px()
# Box/font keys = Excel column names (PARTY_TYPE = PARTY + TYPE text).
#
# Box format: (left, top, right, bottom)  — design-canvas pixels.
# To nudge text, change BOTH edges on that axis so the box size stays the same:
#   UP    → decrease top & bottom   e.g. (L, T-20, R, B-20)
#   DOWN  → increase top & bottom   e.g. (L, T+20, R, B+20)
#   LEFT  → decrease left & right   e.g. (L-20, T, R-20, B)
#   RIGHT → increase left & right   e.g. (L+20, T, R+20, B)
# Smaller step = finer move (try 10–30). Re-run: python generate_cards.py --prototype --no-stickers
# Font size for a field is under the same key inside "fonts" (pt on the design canvas).
_LAYOUT_DESIGNS: dict[str, dict[str, Any]] = {
    "ally": {
        "IMAGE": (52, 198, 1452, 1298),
        "NAME": (116, 66, 904, 136),
        "COST": (1196, 94, 1418, 228),
        "SUBTYPE": (518, 1315, 972, 1381),
        "EFFECT": (268, 1467, 1186, 1685),
        "VC": (92, 1839, 282, 1984),
        "PARTY_TYPE": (395, 1870, 1120, 1928),
        "PARTY": (1220, 1805, 1344, 1933),
        "PARTYABB": (1202, 1986, 1372, 2030),
        "FOOTER_CLEAR": (441, 1961, 776, 2036),
        "SET": (441, 1961, 554, 2005),
        "YEAR": (441, 2014, 760, 2034),
        "ID": (800, 1966, 1055, 1994),
        "CARDNUM": (800, 1992, 1055, 2028),
        "SKILLS_RIGHT": 1402,
        "SKILLS_TOP": 382,
        "SKILLS_GAP": 10,
        "SKILLS_HEIGHT": 58,
        "fonts": {
            "NAME": 10, "COST": 24, "VC": 24, "SUBTYPE": 9,
            "EFFECT": 8, "POWER_COST": 8, "PARTY_TYPE": 8, "PARTYABB": 5.76,
            "SET": 4.6, "YEAR": 2.3,
            "ID": 3.84, "CARDNUM": 4.8,
        },
    },
    "can": {
        "IMAGE": (56, 248, 1444, 1300),
        "NAME": (360, 48, 1190, 142),
        "PARTY_TYPE": (400, 168, 1120, 220),
        "MP": (70, 120, 360, 320),
        "REGEN": (1160, 120, 1450, 320),
        "EFFECT": (140, 1390, 1360, 1760),
        **_FOOTER,
        "fonts": {
            "NAME": 9, "PARTY_TYPE": 6, "MP": 24, "REGEN": 24,
            "EFFECT": 8, "PARTYABB": 6.86, **_FOOTER_FONTS,
        },
    },
    "eff": {
        "IMAGE": (48, 210, 1388, 1290),
        "NAME": (582, 56, 1376, 138),
        "COST": (114, 64, 284, 226),
        "SUBTYPE": (60, 1994, 308, 2040),
        "EFFECT": (140, 1380, 1360, 1710),
        "PARTY_TYPE": (320, 1848, 1140, 1916),
        **_FOOTER,
        "fonts": {
            "NAME": 10, "COST": 30, "SUBTYPE": 6, "EFFECT": 8,
            "PARTY_TYPE": 8, "PARTYABB": 6.86, **_FOOTER_FONTS,
        },
    },
}


def _scale_layout(design: dict[str, Any]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for key, value in design.items():
        if key == "fonts" and isinstance(value, dict):
            out[key] = {fk: pt_px(float(fv)) for fk, fv in value.items()}
        elif isinstance(value, tuple):
            out[key] = tuple(sx(v) for v in value)
        else:
            out[key] = sx(value)
    return out


LAYOUTS = {s: _scale_layout(d) for s, d in _LAYOUT_DESIGNS.items()}


def layout_for(card_type: str) -> dict[str, Any]:
    suffix = TYPE_SUFFIX.get(card_type)
    if not suffix or suffix not in LAYOUTS:
        raise ValueError(f"No layout for card type {card_type!r}")
    return LAYOUTS[suffix]


@dataclass
class CardData:
    id: str
    name: str
    party: str
    type: str
    subtype: str
    mp: Any
    regen: Any
    cost: Any
    vc: Any
    power_cost: Any
    effect: str
    skills: list[str]
    set_name: str
    cardnum: str
    rarity: Any
    set_total: int = 0
    year: int | None = None
    party_abb: str = ""
    image: str = ""

    @property
    def serial(self) -> str:
        return f"# EL-{self.id}"

    @property
    def print_run(self) -> str:
        return f"{str(self.cardnum).zfill(3)}/{self.set_total}"

    @property
    def party_label(self) -> str:
        if self.party_abb:
            return self.party_abb
        return PARTY_SHORT.get(self.party, self.party[:3])


def _num(value: Any) -> int | None:
    if value is None or value in {"-", ""}:
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def _text(value: Any, default: str = "") -> str:
    return default if value is None else str(value).strip()


def parse_skills(raw: Any) -> list[str]:
    if raw is None or str(raw).strip() in {"", "-"}:
        return []
    return [p.upper() for p in re.split(r"[,;/|\s]+", str(raw).strip()) if p and p != "-"]


def _col_map(headers: tuple[Any, ...]) -> dict[str, int]:
    return {_text(h).upper(): i for i, h in enumerate(headers) if h is not None}


def _cell(row: tuple[Any, ...], cols: dict[str, int], *names: str, default: Any = None) -> Any:
    for name in names:
        idx = cols.get(name)
        if idx is not None and idx < len(row):
            return row[idx]
    return default


def load_cards(xlsx_path: Path) -> list[CardData]:
    """Load cards from Excel. Supports new schema (SKILLS1/SKILLS2/SETTOTAL/…) and old (SKILLS)."""
    ws = load_workbook(xlsx_path, data_only=True).active
    rows = ws.iter_rows(values_only=True)
    headers = next(rows)
    cols = _col_map(headers or ())

    cards: list[CardData] = []
    for row in rows:
        if not row or not row[0]:
            continue
        skills: list[str] = []
        seen: set[str] = set()
        for skill in parse_skills(_cell(row, cols, "SKILLS1", "SKILLS")) + parse_skills(
            _cell(row, cols, "SKILLS2")
        ):
            if skill not in seen:
                seen.add(skill)
                skills.append(skill)

        set_total = _num(_cell(row, cols, "SETTOTAL")) or 0
        year = _num(_cell(row, cols, "YEAR"))
        cards.append(
            CardData(
                id=_text(_cell(row, cols, "ID")),
                name=_text(_cell(row, cols, "NAME")),
                party=_text(_cell(row, cols, "PARTY")).upper(),
                type=_text(_cell(row, cols, "TYPE")).upper(),
                subtype=_text(_cell(row, cols, "SUBTYPE")).upper(),
                mp=_cell(row, cols, "MP"),
                regen=_cell(row, cols, "REGEN"),
                cost=_cell(row, cols, "COST"),
                vc=_cell(row, cols, "VC"),
                power_cost=_cell(row, cols, "POWER COST"),
                effect=_text(_cell(row, cols, "EFFECT")),
                skills=skills,
                set_name=_text(_cell(row, cols, "SET"), "BASE").upper() or "BASE",
                cardnum=_text(_cell(row, cols, "CARDNUM")),
                rarity=_cell(row, cols, "RARITY"),
                set_total=set_total,
                year=year,
                party_abb=_text(_cell(row, cols, "PARTYABB")).upper(),
                image=_text(_cell(row, cols, "IMAGE")),
            )
        )

    fallback_total = len(cards)
    for card in cards:
        if not card.set_total:
            card.set_total = fallback_total
    return cards


def template_path(card: CardData) -> Path:
    prefix = PARTY_TEMPLATE_PREFIX.get(card.party)
    suffix = TYPE_SUFFIX.get(card.type)
    if not prefix or not suffix:
        raise ValueError(f"No template mapping for {card.party} {card.type}")
    path = TEMPLATES / f"{prefix}{suffix}.png"
    if not path.exists():
        raise FileNotFoundError(f"Missing template: {path}")
    return path


def load_font(size: int, bold: bool = True) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    # Prefer fonts/ shipped with the project so clients can run on any machine
    bundled = FONT_BOLD if bold else FONT_REGULAR
    for path in (
        str(bundled),
        str(FONT_REGULAR),  # bold missing → regular
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Helvetica.ttc",
        "/Library/Fonts/Arial Unicode.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
        "/System/Library/Fonts/SFNS.ttf",
    ):
        try:
            return ImageFont.truetype(path, size=size, index=0)
        except OSError:
            continue
    return ImageFont.load_default()


def layout_font(layout: dict[str, Any], key: str, *, bold: bool = False) -> ImageFont.ImageFont:
    return load_font(layout["fonts"][key], bold=bold)


def fit_font_to_width(
    text: str, max_width: int, start_px: int, *, bold: bool = False, min_px: int | None = None
) -> ImageFont.ImageFont:
    floor = max(1, min_px if min_px is not None else max(1, int(round(start_px * 0.72))))
    size = start_px
    font = load_font(size, bold=bold)
    probe = ImageDraw.Draw(Image.new("RGB", (1, 1)))
    while size > floor and probe.textlength(text, font=font) > max_width:
        size -= 1
        font = load_font(size, bold=bold)
    return font


def fit_cover(image: Image.Image, box: tuple[int, int, int, int]) -> Image.Image:
    left, top, right, bottom = box
    tw, th = right - left, bottom - top
    src = image.convert("RGBA")
    sw, sh = src.size
    scale = max(tw / sw, th / sh)
    nw, nh = max(1, int(sw * scale)), max(1, int(sh * scale))
    resized = src.resize((nw, nh), Image.Resampling.LANCZOS)
    x, y = (nw - tw) // 2, (nh - th) // 2
    return resized.crop((x, y, x + tw, y + th))


def trim_skill_badge(badge: Image.Image) -> Image.Image:
    """Crop wide skill assets down to the colored pill content."""
    img = badge.convert("RGBA")
    w, h = img.size
    px = img.load()

    def is_content(x: int, y: int) -> bool:
        r, g, b, a = px[x, y]
        return a >= 20 and ((r + g + b) > 40 or a > 200)

    left, right = 0, w - 1
    for x in range(w):
        if any(is_content(x, y) for y in range(h)):
            left = x
            break
    for x in range(w - 1, -1, -1):
        if any(is_content(x, y) for y in range(h)):
            right = x
            break

    dense_right = left
    for x in range(left, right + 1):
        if sum(1 for y in range(h) if is_content(x, y) and sum(px[x, y][:3]) < 700) > h * 0.15:
            dense_right = x
    right = min(right, dense_right + int(h * 0.55))
    return img.crop((max(0, left - 2), 0, min(w, right + 1), h))


def paste_skills(card_img: Image.Image, skills: list[str], layout: dict[str, Any]) -> None:
    if not skills:
        return
    x_right, y = layout["SKILLS_RIGHT"], layout["SKILLS_TOP"]
    target_h, gap = layout["SKILLS_HEIGHT"], layout["SKILLS_GAP"]
    for skill in skills:
        path = ICONS / f"{skill}.png"
        if not path.exists():
            continue
        badge = trim_skill_badge(Image.open(path))
        nw = max(1, int(badge.width * (target_h / badge.height)))
        badge = badge.resize((nw, target_h), Image.Resampling.LANCZOS)
        card_img.alpha_composite(badge, (x_right - nw, y))
        y += target_h + gap


def draw_text(
    draw: ImageDraw.ImageDraw,
    text: str,
    box: tuple[int, int, int, int],
    font: ImageFont.ImageFont,
    *,
    align: str = "center",
    fill=(255, 255, 255, 255),
    stroke_width: int = 0,
) -> None:
    left, top, right, bottom = box
    bbox = draw.textbbox((0, 0), text, font=font, stroke_width=stroke_width)
    tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
    y = top + (bottom - top - th) // 2 - bbox[1]
    if align == "left":
        x = left
    elif align == "right":
        x = max(right - tw - bbox[0], left - bbox[0])
    else:
        x = left + (right - left - tw) // 2 - bbox[0]
    kw: dict[str, Any] = {"font": font, "fill": fill}
    if stroke_width:
        kw.update(stroke_width=stroke_width, stroke_fill=(0, 0, 0, 180))
    draw.text((x, y), text, **kw)


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont, max_width: int) -> list[str]:
    words = text.split()
    if not words:
        return []
    lines: list[str] = []
    current = words[0]
    for word in words[1:]:
        trial = f"{current} {word}"
        if draw.textbbox((0, 0), trial, font=font)[2] <= max_width:
            current = trial
        else:
            lines.append(current)
            current = word
    lines.append(current)
    return lines


def draw_effect_text(draw: ImageDraw.ImageDraw, card: CardData, layout: dict[str, Any]) -> None:
    left, top, right, bottom = layout["EFFECT"]
    effect = card.effect.rstrip(". ").upper()
    if effect and not effect.endswith("."):
        effect += "."

    power_cost = _num(card.power_cost)
    fonts = layout["fonts"]
    effect_px = fonts["EFFECT"]
    if len(effect) >= 60:
        effect_px = max(1, int(round(effect_px * 0.85)))
    font = load_font(effect_px)
    circle_font = load_font(fonts.get("POWER_COST", fonts["EFFECT"]))

    prefix, x, y = "", left + sx(20), top + sx(30)
    if card.type == "ALLY" and power_cost is not None:
        radius = sx(28)
        cx, cy = x + radius, y + radius + sx(4)
        draw.ellipse((cx - radius, cy - radius, cx + radius, cy + radius), fill=(255, 255, 255, 255))
        draw_text(draw, str(power_cost), (cx - radius, cy - radius, cx + radius, cy + radius), circle_font, fill=(30, 40, 60, 255))
        prefix, x = " : ", cx + radius + sx(8)

    body = f"{prefix}{effect}" if prefix else effect
    max_width = right - x - sx(10)
    lines = wrap_text(draw, body, font, max_width)
    line_h = font.size + sx(10)
    limit = 6
    if card.type != "ALLY":
        avail = max(line_h, bottom - y)
        floor = max(1, int(round(fonts["EFFECT"] * 0.68)))
        while len(lines) * line_h > avail and effect_px > floor:
            effect_px -= 1
            font = load_font(effect_px)
            lines = wrap_text(draw, body, font, max_width)
            line_h = font.size + sx(10)
        limit = max(1, avail // max(1, line_h))
    for i, line in enumerate(lines[:limit]):
        draw.text((x if i == 0 else left + sx(20), y + i * line_h), line, font=font, fill=(255, 255, 255, 255))


def draw_subtype_pill(draw: ImageDraw.ImageDraw, subtype: str, layout: dict[str, Any]) -> None:
    if subtype and subtype != "-":
        draw_text(draw, subtype, layout["SUBTYPE"], layout_font(layout, "SUBTYPE", bold=True))


def _knockout_black(icon: Image.Image, threshold: int = 30) -> Image.Image:
    icon = icon.convert("RGBA")
    pixels = icon.load()
    for y in range(icon.height):
        for x in range(icon.width):
            r, g, b, a = pixels[x, y]
            if r < threshold and g < threshold and b < threshold:
                pixels[x, y] = (r, g, b, 0)
    return icon


def place_party_emblem(
    card_img: Image.Image, draw: ImageDraw.ImageDraw, card: CardData, layout: dict[str, Any]
) -> None:
    left, top, right, bottom = layout["PARTY"]
    icon_path = ICONS / f"{card.party}.png"
    if icon_path.exists():
        size = int((right - left) * 0.95)
        icon = _knockout_black(Image.open(icon_path)).resize((size, size), Image.Resampling.LANCZOS)
        card_img.alpha_composite(icon, (left + (right - left - size) // 2, top + (bottom - top - size) // 2))
    draw_text(draw, card.party_label, layout["PARTYABB"], layout_font(layout, "PARTYABB", bold=True))


def _paste_rarity_stars(card_img: Image.Image, origin: tuple[int, int], count: int, size: int, step: int) -> None:
    star_path = ICONS / "rarity.png"
    if not star_path.exists():
        return
    x, y = origin
    star = _knockout_black(Image.open(star_path)).resize((size, size), Image.Resampling.LANCZOS)
    for i in range(count):
        card_img.alpha_composite(star, (x + i * step, y))


def draw_footer(
    card_img: Image.Image, draw: ImageDraw.ImageDraw, card: CardData, layout: dict[str, Any]
) -> None:
    set_name = card.set_name.strip() or "BASE"
    clear = layout["FOOTER_CLEAR"]
    sample = card_img.getpixel((clear[0] - 8, (clear[1] + clear[3]) // 2))
    if isinstance(sample, int):
        panel_fill = (88, 109, 146, 255)
    else:
        panel_fill = sample if len(sample) == 4 else (*sample[:3], 255)
    draw.rectangle(clear, fill=panel_fill)

    badge = layout["SET"]
    draw.rectangle(badge, fill=(255, 255, 255, 255))
    draw_text(draw, set_name, badge, layout_font(layout, "SET", bold=True), fill=(18, 22, 30, 255))

    star_h = sx(28)
    rarity = _num(card.rarity)
    _paste_rarity_stars(
        card_img,
        (badge[2] + sx(12), badge[1] + (badge[3] - badge[1] - star_h) // 2),
        rarity if rarity is not None and rarity > 0 else 5,
        star_h,
        sx(32),
    )

    year = card.year if card.year else 2026
    draw_text(
        draw, f"©{year} TINKER GAMES & HOBBY.", layout["YEAR"],
        layout_font(layout, "YEAR", bold=False), align="left",
    )
    draw_text(draw, card.serial, layout["ID"], layout_font(layout, "ID", bold=False))
    draw_text(draw, card.print_run, layout["CARDNUM"], layout_font(layout, "CARDNUM", bold=True))


def avatar_for(card: CardData, source: Path = AVATAR_DIR) -> Path:
    if source.is_file():
        return source
    directory = source if source.is_dir() else AVATAR_DIR
    # Prefer ID (e.g. 0000001.png) for now; fall back to CARDNUM / IMAGE / default
    candidates = [
        f"{card.id}.png" if card.id else "",
        f"{str(card.cardnum).zfill(3)}.png" if card.cardnum else "",
        f"{card.image}.png" if card.image and card.image != card.id else "",
    ]
    for name in candidates:
        if name:
            path = directory / name
            if path.is_file():
                return path
    fallback = directory / "default-img.png"
    return fallback if fallback.is_file() else DEFAULT_AVATAR


def render_card(card: CardData, avatar_path: Path) -> Image.Image:
    layout = layout_for(card.type)
    template = Image.open(template_path(card)).convert("RGBA")
    if template.size != (CARD_W, CARD_H):
        template = template.resize((CARD_W, CARD_H), Image.Resampling.LANCZOS)

    canvas = Image.new("RGBA", (CARD_W, CARD_H), (20, 24, 32, 255))
    portrait_box = layout["IMAGE"]
    if avatar_path.exists():
        canvas.paste(fit_cover(Image.open(avatar_path), portrait_box), (portrait_box[0], portrait_box[1]))
    canvas.alpha_composite(template)
    draw = ImageDraw.Draw(canvas)
    alliance = f"{card.party} {card.type}"
    stroke = max(1, sx(1))

    if card.type == "ALLY":
        draw_text(
            draw, card.name.upper(), layout["NAME"], layout_font(layout, "NAME", bold=True),
            align="left", stroke_width=stroke,
        )
        power, votes = _num(card.cost), _num(card.vc)
        if power is not None:
            draw_text(draw, str(power), layout["COST"], layout_font(layout, "COST"))
        if votes is not None:
            draw_text(draw, str(votes), layout["VC"], layout_font(layout, "VC"))
        draw_text(draw, alliance, layout["PARTY_TYPE"], layout_font(layout, "PARTY_TYPE", bold=True))
        draw_subtype_pill(draw, card.subtype, layout)
        place_party_emblem(canvas, draw, card, layout)
        paste_skills(canvas, card.skills, layout)
        draw_effect_text(draw, card, layout)
    elif card.type == "CANDIDATE":
        mp, regen = _num(card.mp), _num(card.regen)
        if mp is not None:
            draw_text(draw, str(mp), layout["MP"], layout_font(layout, "MP"))
        if regen is not None:
            draw_text(draw, str(regen), layout["REGEN"], layout_font(layout, "REGEN"))
        draw_text(draw, card.name.upper(), layout["NAME"], layout_font(layout, "NAME", bold=True))
        draw_text(draw, alliance, layout["PARTY_TYPE"], layout_font(layout, "PARTY_TYPE", bold=True))
        draw_effect_text(draw, card, layout)
        place_party_emblem(canvas, draw, card, layout)
    else:  # EFFECT
        name, name_box = card.name.upper(), layout["NAME"]
        name_font = fit_font_to_width(name, name_box[2] - name_box[0], layout["fonts"]["NAME"], bold=True)
        draw_text(draw, name, name_box, name_font, align="right", stroke_width=stroke)
        cost = _num(card.cost)
        if cost is not None:
            draw_text(draw, str(cost), layout["COST"], layout_font(layout, "COST"))
        draw_subtype_pill(draw, card.subtype, layout)
        draw_text(draw, alliance, layout["PARTY_TYPE"], layout_font(layout, "PARTY_TYPE", bold=True))
        place_party_emblem(canvas, draw, card, layout)
        draw_effect_text(draw, card, layout)

    draw_footer(canvas, draw, card, layout)
    return canvas.convert("RGB")


def draw_layout_debug(img: Image.Image, layout: dict[str, Any]) -> Image.Image:
    out = img.convert("RGBA")
    overlay = Image.new("RGBA", out.size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay)
    font = load_font(sx(28), bold=True)
    colors = [
        (255, 64, 64, 220), (64, 220, 64, 220), (64, 140, 255, 220),
        (255, 200, 40, 220), (255, 80, 220, 220), (40, 230, 230, 220), (255, 140, 40, 220),
    ]
    i = 0
    for key, value in layout.items():
        if key == "fonts":
            continue
        color = colors[i % len(colors)]
        i += 1
        if isinstance(value, tuple) and len(value) == 4:
            left, top, right, bottom = value
            draw.rectangle((left, top, right, bottom), outline=color, width=max(1, sx(4)))
            label = f"{key}  ({left},{top})-({right},{bottom})"
            draw.rectangle((left, max(0, top - sx(34)), left + sx(8) * len(label), top), fill=(*color[:3], 180))
            draw.text((left + sx(4), max(0, top - sx(32))), label, font=font, fill=(0, 0, 0, 255))
        elif key == "SKILLS_RIGHT":
            x = int(value)
            draw.line((x, 0, x, CARD_H), fill=color, width=max(1, sx(3)))
            draw.text((x - sx(220), sx(20)), f"SKILLS_RIGHT={x}", font=font, fill=color)
        elif key == "SKILLS_TOP":
            y = int(value)
            draw.line((0, y, CARD_W, y), fill=color, width=max(1, sx(3)))
            draw.text((sx(20), y + sx(6)), f"SKILLS_TOP={y}", font=font, fill=color)

    ruler = load_font(sx(22), bold=False)
    step = sx(100)
    for x in range(0, CARD_W + 1, step):
        draw.line((x, 0, x, sx(18)), fill=(255, 255, 255, 200), width=1)
        draw.text((x + sx(2), sx(20)), str(x), font=ruler, fill=(255, 255, 255, 200))
    for y in range(0, CARD_H + 1, step):
        draw.line((0, y, sx(18), y), fill=(255, 255, 255, 200), width=1)
        draw.text((sx(22), y), str(y), font=ruler, fill=(255, 255, 255, 200))
    return Image.alpha_composite(out, overlay).convert("RGB")


def save_pdf(
    img: Image.Image,
    pdf_path: Path,
    *,
    dpi: int,
    x: float,
    y: float,
    width: float,
    height: float,
    preserve_aspect: bool = True,
) -> None:
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    c = pdf_canvas.Canvas(str(pdf_path), pagesize=letter)
    tmp = pdf_path.with_suffix(".tmp.png")
    try:
        img.save(tmp, "PNG", dpi=(dpi, dpi))
        c.drawImage(str(tmp), x, y, width=width, height=height, preserveAspectRatio=preserve_aspect, mask="auto")
        c.showPage()
        c.save()
    finally:
        tmp.unlink(missing_ok=True)


def save_card_pdf(img: Image.Image, pdf_path: Path) -> None:
    page_w, page_h = letter
    cw, ch = CARD_INCH_W * inch, CARD_INCH_H * inch
    save_pdf(img, pdf_path, dpi=DPI, x=(page_w - cw) / 2, y=(page_h - ch) / 2, width=cw, height=ch)


def compose_sticker_sheet(card_images: list[Image.Image], template_path: Path = STICKER_TEMPLATE) -> Image.Image:
    if not template_path.exists():
        raise FileNotFoundError(f"Missing sticker template: {template_path}")
    sheet = Image.open(template_path).convert("RGB")
    for img, (x, y) in zip(card_images[:STICKER_PER_SHEET], STICKER_SLOTS):
        sticker = img.convert("RGB").resize((STICKER_CARD_W_PX, STICKER_CARD_H_PX), Image.Resampling.LANCZOS)
        ox = x + max(0, (709 - sticker.width) // 2)
        oy = y + max(0, (1005 - sticker.height) // 2)
        sheet.paste(sticker, (ox, oy))
    return sheet


def save_sticker_sheets(
    rendered: list[tuple[str, Image.Image]],
    out_dir: Path,
    template_path: Path = STICKER_TEMPLATE,
) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    page_w, page_h = letter
    total_sheets = max(1, (len(rendered) + STICKER_PER_SHEET - 1) // STICKER_PER_SHEET)

    for sheet_idx in range(total_sheets):
        chunk = rendered[sheet_idx * STICKER_PER_SHEET : (sheet_idx + 1) * STICKER_PER_SHEET]
        if not chunk:
            break
        sheet = compose_sticker_sheet([img for _, img in chunk], template_path)
        stem = f"sticker_sheet_{sheet_idx + 1:03d}"
        png_path, pdf_path = out_dir / f"{stem}.png", out_dir / f"{stem}.pdf"
        sheet.save(png_path, "PNG", optimize=True, dpi=(STICKER_DPI, STICKER_DPI))
        save_pdf(sheet, pdf_path, dpi=STICKER_DPI, x=0, y=0, width=page_w, height=page_h, preserve_aspect=False)
        written.append(pdf_path)
        names = ", ".join(name for name, _ in chunk)
        print(f"Wrote {pdf_path} ({len(chunk)} cards: {names})")
        print(f"Wrote {png_path}")
    return written


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate ELECTOR cards from Excel + templates")
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--avatar", type=Path, default=AVATAR_DIR, help="Avatar folder or one portrait for every card")
    parser.add_argument("--out", type=Path, default=OUTPUT_DIR)
    parser.add_argument("--ids", nargs="*", help="Optional card IDs (default: all)")
    parser.add_argument("--limit", type=int, default=None, help="Only generate first N matching cards")
    parser.add_argument("--prototype", action="store_true", help="Generate Billy Snow only")
    parser.add_argument("--debug-layout", action="store_true", help="Overlay LAYOUT boxes on PNGs")
    parser.add_argument("--no-stickers", action="store_true", help="Skip sticker-sheet PDFs")
    args = parser.parse_args()

    cards = load_cards(args.xlsx)
    if args.prototype:
        cards = [c for c in cards if c.name.upper() == "BILLY SNOW"][:1]
    elif args.ids:
        wanted = {i.lstrip("0") or "0" for i in args.ids} | set(args.ids)
        cards = [c for c in cards if c.id in wanted or c.id.lstrip("0") in wanted]
    if args.limit is not None:
        cards = cards[: args.limit]
    if not cards:
        raise SystemExit("No cards matched the selection.")

    png_dir, pdf_dir, sticker_dir = args.out / "png", args.out / "pdfs", args.out / "stickers"
    for d in (args.out, png_dir, pdf_dir):
        d.mkdir(parents=True, exist_ok=True)

    rendered: list[tuple[str, Image.Image]] = []
    for card in cards:
        avatar_path = avatar_for(card, args.avatar)
        img = render_card(card, avatar_path)
        if args.debug_layout:
            img = draw_layout_debug(img, layout_for(card.type))
        stem = f"{card.cardnum.zfill(3)}_{card.name.replace(' ', '_').upper()}"
        png_path, pdf_path = png_dir / f"{stem}.png", pdf_dir / f"{stem}.pdf"
        img.save(png_path, "PNG", optimize=True, dpi=(DPI, DPI))
        save_card_pdf(img, pdf_path)
        rendered.append((stem, img))
        print(f"Wrote {png_path} (avatar: {avatar_path.name})")
        print(f"Wrote {pdf_path}")

    if not args.no_stickers:
        sheets = save_sticker_sheets(rendered, sticker_dir)
        print(f"Done: {len(cards)} card(s), {len(sheets)} sticker sheet(s) → {sticker_dir}")
    else:
        print(f"Done: {len(cards)} card(s) → {png_dir} and {pdf_dir}")


if __name__ == "__main__":
    main()
